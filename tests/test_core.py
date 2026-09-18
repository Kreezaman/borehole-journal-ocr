from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
import sys

import numpy as np
from openpyxl import load_workbook
from PIL import Image

from app.core.alignment import PageAligner
from app.core.export_xlsx import export_project
from app.core.project_io import load_project, save_project
from app.models import HeaderData, JournalProject, JournalRow, ProjectPage, RecognizedPage
from app.ocr.cloud_gemini import GeminiCloudOCR
from app.ocr.local_paddle import PaddleLocalOCR
from app.ui.workers import SequentialWorker


ROOT = Path(__file__).resolve().parents[1]


def test_frozen_entrypoint_diverts_multiprocessing_before_gui_import() -> None:
    source = (ROOT / "main.py").read_text(encoding="utf-8")
    assert source.index("multiprocessing.freeze_support()") < source.index("from app.ui.main_window import run_app")


def test_alignment_of_template() -> None:
    template = np.asarray(Image.open(ROOT / "app/assets/template.png").convert("RGB"))
    result = PageAligner(template).align(template)
    assert result.image.shape[:2] == (template.shape[0] * 2, template.shape[1] * 2)
    assert result.quality >= 0.8


def test_project_roundtrip(tmp_path: Path) -> None:
    project = JournalProject(
        source_path="sample.pdf",
        pages=[
            ProjectPage(
                page_index=0,
                source_name="sample.pdf — стр. 1",
                source_path="sample.pdf",
                data=RecognizedPage(
                    header=HeaderData(borehole_no="17"),
                    rows=[JournalRow(depth_from="0", depth_to="0,1", description="ПРС")],
                ),
            )
        ],
    )
    output = tmp_path / "sample.bjproj"
    save_project(project, output)
    restored = load_project(output)
    assert restored.pages[0].data.header.borehole_no == "17"
    assert restored.pages[0].data.rows[0].description == "ПРС"


def test_excel_export(tmp_path: Path) -> None:
    project = JournalProject(
        pages=[
            ProjectPage(
                page_index=0,
                source_name="one",
                data=RecognizedPage(
                    header=HeaderData(borehole_no="18А"),
                    rows=[JournalRow(depth_from="0", depth_to="2,2", description="Песок мелкий")],
                ),
            )
        ]
    )
    output = tmp_path / "result.xlsx"
    export_project(project, output)
    workbook = load_workbook(output)
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet["B3"].value == "18А"
    assert sheet["E8"].value == "Песок мелкий"


def test_local_ocr_disables_onednn(monkeypatch) -> None:
    received: dict = {}

    class FakePaddleOCR:
        def __init__(self, **kwargs):
            received.update(kwargs)

    monkeypatch.setitem(sys.modules, "paddleocr", SimpleNamespace(PaddleOCR=FakePaddleOCR))
    PaddleLocalOCR()
    assert received["enable_mkldnn"] is False
    assert received["device"] == "cpu"
    assert received["text_recognition_model_name"] == "cyrillic_PP-OCRv5_mobile_rec"


def test_local_worker_forces_native_process_exit() -> None:
    source = (ROOT / "app/ocr/worker_process.py").read_text(encoding="utf-8")
    assert "temporary_output.replace(output_path)" in source
    assert "os._exit(exit_code)" in source


def test_gemini_ocr_uses_inline_image_and_structured_output() -> None:
    received: dict = {}
    progress: list[str] = []

    class FakeInteractions:
        def create(self, **kwargs):
            received.update(kwargs)
            return SimpleNamespace(
                output_text=_gemini_result_json()
            )

    client = SimpleNamespace(interactions=FakeInteractions())
    image = np.full((80, 120, 3), 255, dtype=np.uint8)
    result = GeminiCloudOCR("test-key", client=client).recognize(image, progress.append)

    assert received["model"] == "gemini-3.1-flash-lite"
    images = [part for part in received["input"] if part["type"] == "image"]
    assert len(images) == 5
    assert images[0]["resolution"] == "medium"
    assert all(image["resolution"] == "high" for image in images[1:])
    assert all(image["mime_type"] == "image/jpeg" for image in images)
    assert received["generation_config"]["thinking_level"] == "minimal"
    assert received["response_format"]["mime_type"] == "application/json"
    assert result.header.borehole_no == "2"
    assert result.rows[0].description == "ПРС"
    assert len(result.rows) == 19
    assert result.recognition_mode == "gemini"
    schema = received["response_format"]["schema"]
    assert "header" in schema["required"]
    assert progress[0].startswith("Этап 1 из 5")
    assert progress[-1].startswith("Этап 5 из 5")


def test_gemini_retries_a_temporary_rate_limit_with_visible_countdown() -> None:
    calls = 0
    progress: list[str] = []
    sleeps: list[float] = []

    class RateLimitError(Exception):
        code = 429
        response = SimpleNamespace(headers={"retry-after": "1"})
        details = {"error": {"status": "RESOURCE_EXHAUSTED"}}

    class FakeInteractions:
        def create(self, **_kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                raise RateLimitError("429 RESOURCE_EXHAUSTED")
            return SimpleNamespace(output_text=_gemini_result_json())

    client = SimpleNamespace(interactions=FakeInteractions())
    image = np.full((80, 120, 3), 255, dtype=np.uint8)
    GeminiCloudOCR("test-key", client=client, sleep_fn=sleeps.append).recognize(
        image, progress.append
    )

    assert calls == 2
    assert sleeps == [1]
    assert any("Повтор через 1 сек" in message for message in progress)


def test_gemini_rechecks_only_a_suspicious_cell() -> None:
    requests: list[dict] = []

    class FakeInteractions:
        def create(self, **kwargs):
            requests.append(kwargs)
            if len(requests) == 1:
                return SimpleNamespace(
                    output_text=_gemini_result_json(
                        depth_to="0,7",
                        warnings=[
                            {
                                "field": "rows[0].depth_to",
                                "reason": "неясно 7 или 1",
                                "alternative": "0,1",
                            }
                        ],
                    )
                )
            return SimpleNamespace(
                output_text=json.dumps(
                    {
                        "results": [
                            {
                                "field": "rows[0].depth_to",
                                "value": "0,1",
                                "uncertain": False,
                                "reason": "",
                            }
                        ]
                    },
                    ensure_ascii=False,
                )
            )

    client = SimpleNamespace(interactions=FakeInteractions())
    image = np.full((1790, 2586, 3), 255, dtype=np.uint8)
    result = GeminiCloudOCR("test-key", client=client).recognize(image)

    assert len(requests) == 2
    review_images = [part for part in requests[1]["input"] if part["type"] == "image"]
    assert len(review_images) == 1
    assert result.rows[0].depth_to == "0,1"
    assert result.ocr_warnings == []


def test_project_still_accepts_legacy_cloud_mode() -> None:
    assert RecognizedPage(recognition_mode="cloud").recognition_mode == "cloud"


def test_sequential_worker_processes_items_in_order() -> None:
    started: list[tuple[int, int, int]] = []
    results: list[tuple[int, int, int, int]] = []
    finished: list[tuple[int, bool]] = []
    cleaned: list[bool] = []

    worker = SequentialWorker(
        [3, 5, 8],
        lambda item, progress: (progress(f"обработка {item}"), item * 2)[1],
        cleanup=lambda: cleaned.append(True),
    )
    worker.signals.item_started.connect(
        lambda item, position, total: started.append((item, position, total))
    )
    worker.signals.item_result.connect(
        lambda item, result, position, total: results.append(
            (item, result, position, total)
        )
    )
    worker.signals.finished.connect(
        lambda completed, stopped: finished.append((completed, stopped))
    )

    worker.run()

    assert started == [(3, 1, 3), (5, 2, 3), (8, 3, 3)]
    assert results == [(3, 6, 1, 3), (5, 10, 2, 3), (8, 16, 3, 3)]
    assert finished == [(3, False)]
    assert cleaned == [True]


def test_sequential_worker_stops_between_items() -> None:
    processed: list[int] = []
    finished: list[tuple[int, bool]] = []
    worker = SequentialWorker([0, 1, 2], lambda item, _progress: item)

    def accept_result(item, _result, _position, _total) -> None:
        processed.append(item)
        worker.request_stop()

    worker.signals.item_result.connect(accept_result)
    worker.signals.finished.connect(
        lambda completed, stopped: finished.append((completed, stopped))
    )

    worker.run()

    assert processed == [0]
    assert finished == [(1, True)]


def test_sequential_worker_keeps_completed_results_before_error() -> None:
    results: list[int] = []
    errors: list[tuple[int, str, int, int]] = []
    finished: list[tuple[int, bool]] = []

    def process(item: int, _progress) -> int:
        if item == 2:
            raise RuntimeError("quota")
        return item * 10

    worker = SequentialWorker([1, 2, 3], process)
    worker.signals.item_result.connect(
        lambda _item, result, _position, _total: results.append(result)
    )
    worker.signals.error.connect(
        lambda item, message, position, total: errors.append(
            (item, message, position, total)
        )
    )
    worker.signals.finished.connect(
        lambda completed, stopped: finished.append((completed, stopped))
    )

    worker.run()

    assert results == [10]
    assert errors == [(2, "RuntimeError: quota", 2, 3)]
    assert finished == [(1, False)]


def _gemini_result_json(
    *,
    depth_to: str = "0,1",
    warnings: list[dict] | None = None,
) -> str:
    return json.dumps(
        {
            "header": {
                "borehole_no": "2",
                "drilling_start": "",
                "drilling_end": "",
                "object_name": "",
                "coordinate_n": "",
                "coordinate_e": "",
            },
            "rows": [
                {
                    "layer_no": "1",
                    "depth_from": "0",
                    "depth_to": depth_to,
                    "frozen_interval": "",
                    "description": "ПРС",
                    "sample_no": "",
                    "sample_depth": "",
                    "water": "",
                }
            ],
            "footer": {
                "groundwater_exposed": "",
                "groundwater_stabilized": "",
                "specific_odor": "",
                "drilling_rig": "",
                "drilling_diameter": "",
                "drilling_depth": "10",
                "drill_master": "",
                "engineer_geologist": "",
                "sketch_notes": "",
            },
            "page_notes": "",
            "warnings": warnings or [],
        },
        ensure_ascii=False,
    )
