from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import JournalProject


HEADERS = [
    "№ слоя",
    "Глубина от, м",
    "Глубина до, м",
    "Мерзлота / мощность",
    "Описание грунта",
    "№ образца",
    "Глубина отбора, м",
    "Воды",
]


def export_project(project: JournalProject, output_path: str | Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    thin = Side(style="thin", color="8A98A8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="244968")
    section_fill = PatternFill("solid", fgColor="DCEAF3")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")

    for page in project.pages:
        data = page.data
        sheet = workbook.create_sheet(_safe_sheet_name(f"Скв {data.header.borehole_no or page.page_index + 1}"))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A8"

        sheet.merge_cells("A1:H1")
        sheet["A1"] = "БУРОВОЙ ЖУРНАЛ"
        sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        sheet["A1"].fill = title_fill
        sheet["A1"].alignment = Alignment(horizontal="center")

        header_rows = [
            ("№ скважины", data.header.borehole_no, "Дата начала", data.header.drilling_start),
            ("Объект", data.header.object_name, "Дата окончания", data.header.drilling_end),
            ("Координата N", data.header.coordinate_n, "Координата E", data.header.coordinate_e),
        ]
        for row_no, (a, b, c, d) in enumerate(header_rows, start=3):
            sheet.cell(row_no, 1, a).font = Font(bold=True)
            sheet.merge_cells(start_row=row_no, start_column=2, end_row=row_no, end_column=4)
            sheet.cell(row_no, 2, b)
            sheet.cell(row_no, 5, c).font = Font(bold=True)
            sheet.merge_cells(start_row=row_no, start_column=6, end_row=row_no, end_column=8)
            sheet.cell(row_no, 6, d)

        for column, value in enumerate(HEADERS, start=1):
            cell = sheet.cell(7, column, value)
            cell.font = Font(bold=True)
            cell.fill = section_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_no, row in enumerate(data.normalized_rows(), start=8):
            values = [
                row.layer_no,
                row.depth_from,
                row.depth_to,
                row.frozen_interval,
                row.description,
                row.sample_no,
                row.sample_depth,
                row.water,
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_no, column, value)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if value.strip() and value.strip() == "?":
                    cell.fill = warning_fill
            sheet.row_dimensions[row_no].height = 42

        footer_row = 29
        footer_values = [
            ("УГВ вскрытый", data.footer.groundwater_exposed),
            ("УГВ установившийся", data.footer.groundwater_stabilized),
            ("Специфический запах", data.footer.specific_odor),
            ("Буровая установка", data.footer.drilling_rig),
            ("Диаметр бурения", data.footer.drilling_diameter),
            ("Глубина бурения", data.footer.drilling_depth),
            ("Буровой мастер", data.footer.drill_master),
            ("Инженер-геолог", data.footer.engineer_geologist),
            ("Абрис / примечания", data.footer.sketch_notes),
        ]
        for offset, (label, value) in enumerate(footer_values):
            r = footer_row + offset
            sheet.cell(r, 1, label).font = Font(bold=True)
            sheet.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            sheet.cell(r, 2, value).alignment = Alignment(wrap_text=True, vertical="top")

        widths = [12, 15, 15, 21, 70, 14, 21, 16]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.auto_filter.ref = "A7:H26"
        sheet.print_title_rows = "1:7"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(output_path)


def _safe_sheet_name(value: str) -> str:
    for symbol in "[]:*?/\\":
        value = value.replace(symbol, "_")
    return value[:31] or "Журнал"
